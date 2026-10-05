# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook('testcases/Dolibarr_TestCases.xlsx')

# 1. Sheet Bug Report
ws_bug = wb['Bug Report']
bug_row = [
    'BUG_001',
    'Redirect sau khi tạo KH chứa placeholder id=__ID__',
    'CRM',
    'TC_CRM_010',
    'Dolibarr 22.0.4, Windows, Chrome, MariaDB 10.6.5',
    '1. Mở /societe/card.php?action=create\n2. Nhập tên KH bất kỳ, chọn loại Khách hàng (customer=1)\n3. Nhấn Lưu (Save)',
    'URL chuyển hướng thực tế: http://localhost/dolibarr/societe/card.php?id=__ID__&socid=109. Placeholder id=__ID__ không được thay thế bằng ID số thực tế.',
    'URL chuyển hướng phải là http://localhost/dolibarr/societe/card.php?socid={id} theo đặc tả UC-02 §4, không chứa placeholder __ID__.',
    'Minor',
    'P2',
    'New',
    'evidence/automation/TC_CRM_010_fail_bug001.png',
    '2026-10-05',
    'Mã nguồn Dolibarr 22.0.4 (htdocs/societe/card.php dòng 204 và 644-648): biến backtopage gán mặc định .../societe/card.php?id=__ID__. Khi redirect, code dùng preg_replace(/--IDFORBACKTOPAGE--/, ...) thay vì __ID__, rồi nối thêm &socid={id}, khiến __ID__ bị giữ nguyên trên URL. Ảnh hưởng: không mất dữ liệu nhưng sai URL chuẩn và có thể gây lỗi khi bookmark hoặc quay lại trang.'
]
for col_idx, val in enumerate(bug_row, 1):
    ws_bug.cell(2, col_idx, val)

# 2. Sheet Test Cases
ws_tc = wb['Test Cases']
header = [ws_tc.cell(1, c).value for c in range(1, 25)]
col_map = {name: idx + 1 for idx, name in enumerate(header) if name}

for r in range(2, ws_tc.max_row + 1):
    t_id = ws_tc.cell(r, 1).value
    if t_id == 'TC_CRM_009':
        ws_tc.cell(r, col_map['Ghi chu'], 'EP ngoài biên: khoảng trắng thuần túy. Expected theo quan sát trong kiểm thử ngày 2026-10-03, chờ giảng viên xác nhận (chưa phải spec chính thức).')
    elif t_id == 'TC_CRM_010':
        ws_tc.cell(r, col_map['Expected'], 'Form lưu thành công; trang chi tiết không thực thi thẻ HTML/script (an toàn XSS); giữ phần nội dung hợp lệ O\'Brien; không có PHP Fatal error. Ký tự <, > không xuất hiện trong tên. Lưu ý: Ký tự ngoặc kép " bị server loại bỏ nhưng test code chưa assert (chờ giảng viên xác nhận).')
        ws_tc.cell(r, col_map['Actual'], 'Lưu KH thành công nhưng redirect về URL sai: societe/card.php?id=__ID__&socid=109 (BUG_001). Server loại bỏ <Test> và ", giữ \' và &, tên hiển thị "O\'Brien & Cong ty 123". An toàn XSS, nhưng test Fail do URL redirect chứa placeholder __ID__. Dấu " chưa được test.')
        ws_tc.cell(r, col_map['Trang thai'], 'Fail')
        ws_tc.cell(r, col_map['Ngay chay'], '2026-10-05')
        ws_tc.cell(r, col_map['Minh chung'], 'evidence/automation/TC_CRM_010_fail_bug001.png')
        ws_tc.cell(r, col_map['Bug ID'], 'BUG_001')
        ws_tc.cell(r, col_map['Ghi chu'], 'Test tự động phát hiện BUG_001 (URL chuyển hướng chứa placeholder id=__ID__). Mục tiêu XSS đạt, phần dấu " chưa assert trong code và chờ giảng viên xác nhận spec.')

# 3. Sheet Summary
ws_sum = wb['Summary']
ws_sum.cell(2, 7, 6) # F-CRM-01 Passed
ws_sum.cell(2, 8, 1) # F-CRM-01 Failed
ws_sum.cell(6, 7, 11) # Tong cong Passed
ws_sum.cell(6, 8, 1)  # Tong cong Failed

# 4. Sheet AI Log
ws_ai = wb['AI Log']
ai_row = [
    8,
    '2026-10-05',
    'Kiểm tra chất lượng / Tối ưu code',
    'Antigravity (Gemini Flash)',
    'Rà soát tính độc lập của Expected TC_CRM_009/010, truy vết nguồn gốc __ID__ trong URL redirect, mở BUG_001 và điều chỉnh phương án kiểm thử XSS',
    'Chấn chỉnh cách xử lý che giấu lỗi bằng NormalizeCustomerUrl; yêu cầu assert URL thô để bắt BUG_001; tách bạch spec chính thức với quan sát kiểm thử; ghi nhận BUG_001 vào Excel',
    'Phát hiện nguyên nhân gốc tại Dolibarr htdocs/societe/card.php:204 & 644-648. Thêm assert URL thô không chứa __ID__ làm fail TC_CRM_010 để bắt lỗi thật (BUG_001). Ghi nhận BUG_001 vào Bug Report, cập nhật CHANGELOG và Excel.',
    'Phải sửa (Rủi ro che giấu lỗi kiểm thử)',
    'Phát hiện AI viết hàm NormalizeCustomerUrl để làm đẹp kết quả khiến URL hợp lệ hóa thay vì bắt lỗi vi phạm spec; yêu cầu AI assert URL thô, ghi nhận bug của ứng dụng thay vì fix ở client.',
    'Mục đích cốt lõi của kiểm thử tự động là phơi bày lỗi của ứng dụng (AUT) đối chiếu với đặc tả, tuyệt đối không được viết mã workaround ở test script để \'làm đẹp\' kết quả hoặc biến lỗi của hệ thống thành test pass.'
]
ws_ai.cell(9, 1, ai_row[0])
for col_idx, val in enumerate(ai_row, 1):
    ws_ai.cell(9, col_idx, val)

wb.save('testcases/Dolibarr_TestCases.xlsx')
print('EXCEL UPDATED SUCCESSFULLY')
