# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook('testcases/Dolibarr_TestCases.xlsx')
ws = wb['Summary']

# Row 2 to 5: F-CRM-01 to F-CRM-04
for r in range(2, 6):
    ws[f'D{r}'] = f"=COUNTIF('Test Cases'!$C:$C, B{r})"
    ws[f'E{r}'] = f'=COUNTIFS(\'Test Cases\'!$C:$C, B{r}, \'Test Cases\'!$H:$H, "Auto")'
    ws[f'F{r}'] = f'=COUNTIFS(\'Test Cases\'!$C:$C, B{r}, \'Test Cases\'!$H:$H, "Manual")'
    ws[f'G{r}'] = f'=COUNTIFS(\'Test Cases\'!$C:$C, B{r}, \'Test Cases\'!$N:$N, "Pass")'
    ws[f'H{r}'] = f'=COUNTIFS(\'Test Cases\'!$C:$C, B{r}, \'Test Cases\'!$N:$N, "Fail")'
    ws[f'I{r}'] = f'=COUNTIFS(\'Test Cases\'!$C:$C, B{r}, \'Test Cases\'!$N:$N, "Not Run")'
    ws[f'J{r}'] = f'=IF(D{r}>0, G{r}/D{r}, 0)'
    ws[f'J{r}'].number_format = '0.0%'

# Row 6: TONG CONG
ws['D6'] = '=SUM(D2:D5)'
ws['E6'] = '=SUM(E2:E5)'
ws['F6'] = '=SUM(F2:F5)'
ws['G6'] = '=SUM(G2:G5)'
ws['H6'] = '=SUM(H2:H5)'
ws['I6'] = '=SUM(I2:I5)'
ws['J6'] = '=IF(D6>0, G6/D6, 0)'
ws['J6'].number_format = '0.0%'

wb.save('testcases/Dolibarr_TestCases.xlsx')
print('SUCCESSFULLY_UPDATED_SUMMARY_FORMULAS')
